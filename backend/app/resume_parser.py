import os
import re
import json
import urllib.request
import urllib.error

def get_pdf_text(filepath):
    try:
        from pypdf import PdfReader
        reader = PdfReader(filepath)
        text = ""
        for page in reader.pages:
            t = page.extract_text()
            if t:
                text += t + "\n"
        return text
    except Exception as e:
        print(f"Error reading PDF: {e}")
        return ""

def get_docx_text(filepath):
    try:
        import docx
        doc = docx.Document(filepath)
        text = []
        for p in doc.paragraphs:
            if p.text.strip():
                text.append(p.text)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        text.append(cell.text)
        return "\n".join(text)
    except Exception as e:
        print(f"Error reading DOCX: {e}")
        return ""

def get_file_text(filepath):
    if not os.path.exists(filepath):
        return ""
    ext = os.path.splitext(filepath)[1].lower()
    if ext == ".pdf":
        return get_pdf_text(filepath)
    elif ext in [".docx", ".doc"]:
        return get_docx_text(filepath)
    else:
        try:
            with open(filepath, "r", encoding="utf-8", errors="ignore") as f:
                return f.read()
        except Exception as e:
            print(f"Error reading raw text: {e}")
            return ""

def extract_heuristics(text, default_name=""):
    lines = [line.strip() for line in text.split("\n") if line.strip()]
    
    # 1. Email
    emails = re.findall(r'[\w\.-]+@[\w\.-]+\.\w+', text)
    email = emails[0] if emails else ""
    
    # 2. Phone
    phones = re.findall(r'(?:\+?\d{1,3}[-.\s]?)?\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}|\b\d{10}\b', text)
    phone = phones[0] if phones else ""
    
    # 3. Name
    name = ""
    for line in lines[:5]:
        if "@" not in line and not any(k in line.lower() for k in ["resume", "cv", "curriculum", "page", "contact", "email", "phone", "profile", "objective"]):
            words = line.split()
            if 1 < len(words) <= 4:
                name = line
                break
    if not name:
        name = default_name
        
    # 4. Location
    location = ""
    for line in lines:
        if any(k in line.lower() for k in ["address:", "location:", "lives in:", "resident of:"]):
            location = re.sub(r'(?i)address:|location:|lives in:|resident of:', '', line).strip()
            break
            
    # 5. Summary / Objective
    summary = ""
    for idx, line in enumerate(lines):
        if any(k == line.lower() for k in ["summary", "objective", "profile", "about me", "professional summary"]):
            summary_lines = []
            for j in range(1, 4):
                if idx + j < len(lines):
                    next_line = lines[idx+j]
                    if any(h in next_line.lower() for h in ["experience", "education", "skills", "projects", "employment"]):
                        break
                    summary_lines.append(next_line)
            summary = " ".join(summary_lines)
            break
            
    # 6. Skills
    skill_keywords = [
        "python", "javascript", "typescript", "html", "css", "react", "node", "sql", "git", "java", "c++",
        "teaching", "management", "communication", "marketing", "leadership", "excel", "social work",
        "writing", "design", "photoshop", "yoga", "art", "music", "event planning", "fundraising",
        "public relations", "translation", "coordination", "tutoring", "ngo", "community organizing",
        "public speaking", "problem solving", "teamwork", "research", "counseling", "first aid"
    ]
    skills = []
    for sk in skill_keywords:
        pattern = r'\b' + re.escape(sk) + r'\b'
        if re.search(pattern, text.lower()):
            skills.append(sk.capitalize() if sk not in ["html", "css", "sql"] else sk.upper())
            
    for idx, line in enumerate(lines):
        if any(k == line.lower() for k in ["skills", "technical skills", "core competencies", "skills & expertise", "expertise"]):
            for j in range(1, 6):
                if idx + j < len(lines):
                    next_line = lines[idx+j]
                    if any(h in next_line.lower() for h in ["experience", "education", "projects", "employment"]):
                        break
                    parts = re.split(r',|•|\||\*|;', next_line)
                    for pt in parts:
                        pt_s = pt.strip()
                        if pt_s and len(pt_s) < 30 and pt_s not in skills:
                            skills.append(pt_s)
            break

    # 7. Education
    education = []
    current_edu = None
    edu_section = False
    for idx, line in enumerate(lines):
        line_lower = line.lower()
        if any(k == line_lower for k in ["education", "academic profile", "qualifications", "academic background", "academic qualifications"]):
            edu_section = True
            continue
        if edu_section:
            if any(h in line_lower for h in ["experience", "skills", "projects", "employment", "languages", "hobbies"]):
                edu_section = False
                break
            
            year_match = re.search(r'\b(19|20)\d{2}\b', line)
            degree_match = any(d in line for d in ["B.Tech", "M.Tech", "B.Sc", "M.Sc", "B.A", "M.A", "MBA", "Ph.D", "B.Com", "M.Com", "School", "University", "College", "Degree", "Diploma", "Btech", "Mtech", "Graduate", "Class X", "Class XII", "10th", "12th"])
            
            if degree_match or year_match:
                if current_edu:
                    education.append(current_edu)
                current_edu = {
                    "degree": line if len(line) < 100 else "Education Detail",
                    "institution": "",
                    "year": year_match.group(0) if year_match else "",
                    "details": ""
                }
                if idx + 1 < len(lines):
                    next_line = lines[idx+1]
                    if not any(h in next_line.lower() for h in ["experience", "skills", "projects", "education"]) and len(next_line) < 120:
                        current_edu["institution"] = next_line
            elif current_edu and len(line) < 150:
                if not current_edu["institution"]:
                    current_edu["institution"] = line
                else:
                    current_edu["details"] = (current_edu["details"] + " " + line).strip()
    if current_edu:
        education.append(current_edu)
        
    # 8. Experience
    professional = []
    current_exp = None
    exp_section = False
    for idx, line in enumerate(lines):
        line_lower = line.lower()
        if any(k == line_lower for k in ["experience", "work history", "employment", "professional experience", "work experience", "employment history"]):
            exp_section = True
            continue
        if exp_section:
            if any(h in line_lower for h in ["education", "skills", "projects", "languages", "hobbies", "interests"]):
                exp_section = False
                break
                
            duration_match = re.search(r'\b(?:19|20)\d{2}\b', line)
            job_match = any(jw in line_lower for jw in ["manager", "developer", "teacher", "engineer", "designer", "coordinator", "volunteer", "intern", "associate", "consultant", "officer", "executive", "lead", "analyst"])
            
            if job_match or duration_match:
                if current_exp:
                    professional.append(current_exp)
                current_exp = {
                    "role": line if len(line) < 80 else "Professional Role",
                    "company": "",
                    "duration": duration_match.group(0) if duration_match else "",
                    "description": ""
                }
                if idx + 1 < len(lines):
                    next_line = lines[idx+1]
                    if not any(h in next_line.lower() for h in ["experience", "education", "skills"]) and len(next_line) < 100:
                        current_exp["company"] = next_line
            elif current_exp:
                if not current_exp["company"] and len(line) < 80:
                    current_exp["company"] = line
                else:
                    current_exp["description"] = (current_exp["description"] + " " + line).strip()
    if current_exp:
        professional.append(current_exp)

    return {
        "personal": {
            "name": name or default_name,
            "email": email,
            "phone": phone,
            "location": location,
            "summary": summary
        },
        "professional": professional,
        "education": education,
        "skills": skills
    }

def analyze_with_gemini(text, api_key):
    url = f"https://generativelanguage.googleapis.com/v1beta/models/gemini-1.5-flash:generateContent?key={api_key}"
    prompt = (
        "You are an expert resume parser and CV analyzer. Your job is to extract structural information from the following resume text. "
        "Provide your output strictly in JSON format. Do not use any markdown formatting, code block wrappers (like ```json), or extra text. "
        "The JSON object must have the following keys:\n"
        "- personal: an object containing 'name', 'email', 'phone', 'location', 'summary'\n"
        "- professional: an array of objects, each containing 'role', 'company', 'duration', 'description'\n"
        "- education: an array of objects, each containing 'degree', 'institution', 'year', 'details'\n"
        "- skills: an array of strings representing individual skills\n\n"
        f"Resume Text:\n{text}"
    )
    
    headers = {"Content-Type": "application/json"}
    data = {
        "contents": [{
            "parts": [{"text": prompt}]
        }],
        "generationConfig": {
            "responseMimeType": "application/json"
        }
    }
    
    req = urllib.request.Request(
        url, 
        data=json.dumps(data).encode("utf-8"), 
        headers=headers, 
        method="POST"
    )
    
    try:
        with urllib.request.urlopen(req, timeout=15) as response:
            res_data = json.loads(response.read().decode("utf-8"))
            content_text = res_data['candidates'][0]['content']['parts'][0]['text']
            return json.loads(content_text)
    except Exception as e:
        print(f"Error calling Gemini API: {e}")
        return None

def parse_resume(filepath, default_name=""):
    text = get_file_text(filepath)
    if not text.strip():
        return {
            "personal": {"name": default_name, "email": "", "phone": "", "location": "", "summary": "No text content found in resume file."},
            "professional": [],
            "education": [],
            "skills": []
        }
    
    api_key = os.getenv("GEMINI_API_KEY")
    if api_key:
        result = analyze_with_gemini(text, api_key)
        if result:
            return result
            
    return extract_heuristics(text, default_name)
